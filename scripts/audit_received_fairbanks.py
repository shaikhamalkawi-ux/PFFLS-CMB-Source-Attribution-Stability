#!/usr/bin/env python3
"""Audit privately received Fairbanks workbooks without redistributing records.

No XLSX is written or formula executed. The hash-locked sources are read twice
(formulas and cached values). Primary matches use exact site/calendar date,
positive measured PM and levoglucosan, and both valid *original* CMB systems.
Native Excel dates and text-date candidates are counted separately. CSVs,
including formula/cached-value evidence, must stay below private/received_fairbanks.
Only an aggregate, identifier-free JSON summary may be written publicly.
"""

from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date, datetime
from decimal import Decimal
import hashlib
from io import BytesIO
import json
from pathlib import Path
from statistics import mean
import warnings
import zipfile


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_ZIP_SHA256 = "e0498220b12ef493bd3d3a16fa27c17becefd8495f0123463a78053cc85d05ac"
EXPECTED_WORKBOOKS = {
    "Fairbanks summary 08-11_cpp.xlsx": "973917d6c504f25980938a8a45da62ce766099734e053a65882c2c05391d8ee7",
    "LevoglucosanResults_final1.xlsx": "1541b1102f49fec9d461fbfb9c3095c13d7b8a23cbce54c81f1744fe7572691a",
}
SITES = {
    "State Building": "state_building",
    "North Pole": "north_pole",
    "Peger Rd": "peger_road",
    "RAMS": "rams",
}
FINE_CF = Decimal("9.01")
OMNI_CF = Decimal("13.27")


def number(value):
    """Finite numeric source value, without converting missing/flags to zero."""
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    parsed = Decimal(str(value))
    return parsed if parsed.is_finite() else None


def value_flag(value):
    if value is None or value == "":
        return "missing"
    numeric = number(value)
    if numeric is not None:
        return "negative" if numeric < 0 else "zero" if numeric == 0 else "numeric"
    text = str(value).strip()
    if text.startswith("#"):
        return "excel_error"
    if text.lower() == "<error":
        return "source_less_than_error_flag"
    if text.startswith(("<", ">", "≤", "≥")):
        return "censored_or_limit_flag"
    return "source_text_flag"


def date_value(value):
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            raise ValueError("Unexpected non-midnight sampling timestamp")
        return value.date().isoformat(), "excel_date"
    if isinstance(value, date):
        return value.isoformat(), "excel_date"
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%m/%d/%Y").date().isoformat(), "text_date"
        except ValueError:
            pass
    return None, "not_date"


def serial(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def cell_evidence(formula_row, cached_row, count):
    """One JSON CSV field prevents formula text becoming executable CSV cells."""
    return json.dumps([
        {"cell": f.coordinate, "source": serial(f.value), "cached": serial(c.value),
         "formula": f.data_type == "f", "number_format": f.number_format}
        for f, c in zip(formula_row[:count], cached_row[:count])
        if f.value is not None or c.value is not None
    ], ensure_ascii=False, sort_keys=True)


def marker_bracket(pm25, levoglucosan, fine=FINE_CF, omni=OMNI_CF):
    mass, levo = number(pm25), number(levoglucosan)
    if mass is None or mass <= 0 or levo is None or levo <= 0:
        return None
    # ng/m3 -> ug/m3, then percent of measured PM2.5; no rounding/clipping.
    share = Decimal("100") * Decimal("0.001") * levo / mass
    return share * fine, share * omni


def interval_metrics(value, low, high):
    value, low, high = map(Decimal, (value, low, high))
    if low > high:
        raise ValueError("Reversed interval is not silently sorted")
    relation = "below" if value < low else "above" if value > high else "inside"
    distance = max(low - value, value - high, Decimal(0))
    midpoint_difference = value - (low + high) / 2
    return {"relation": relation, "interval_distance_pp": distance,
            "midpoint_signed_difference_pp": midpoint_difference,
            "midpoint_absolute_difference_pp": abs(midpoint_difference)}


def _paired_rows(formula_sheet, cached_sheet):
    return zip(formula_sheet.iter_rows(), cached_sheet.iter_rows())


def parse_levoglucosan(formula_book, cached_book):
    rows = []
    for sheet_name, site in SITES.items():
        fs, cs = formula_book[sheet_name], cached_book[sheet_name]
        fine, omni = number(cs["F3"].value), number(cs["G3"].value)
        if (fine, omni) != (FINE_CF, OMNI_CF):
            raise ValueError(f"Unexpected conversion factors in {sheet_name}")
        if (cs["F4"].value, cs["G4"].value) != ("Fine", "OMNI"):
            raise ValueError(f"Unexpected conversion labels in {sheet_name}")
        for fr, cr in _paired_rows(fs, cs):
            day, kind = date_value(cr[0].value)
            if day is None:
                continue
            index = cr[0].row
            pm, levo = cr[2].value, cr[3].value
            bracket = marker_bracket(pm, levo, fine, omni)
            row = {
                "site": site, "date": day, "date_kind": kind,
                "workbook": "LevoglucosanResults_final1.xlsx", "sheet": sheet_name,
                "source_range": f"A{index}:G{index}",
                "filter_id_json": json.dumps(serial(cr[1].value)),
                "filter_id_present": cr[1].value not in (None, ""),
                "pm25_ug_m3": number(pm), "pm25_flag": value_flag(pm),
                "levoglucosan_ng_m3": number(levo), "levoglucosan_flag": value_flag(levo),
                "fine_cf": fine, "omni_cf": omni, "factor_cells": "F3:G4",
                "fine_ws_percent": bracket[0] if bracket else None,
                "omni_ws_percent": bracket[1] if bracket else None,
                "fine_cached_percent": number(cr[5].value),
                "omni_cached_percent": number(cr[6].value),
                "fine_cache_flag": value_flag(cr[5].value),
                "omni_cache_flag": value_flag(cr[6].value),
                "fine_formula_present": fr[5].data_type == "f",
                "omni_formula_present": fr[6].data_type == "f",
                "fine_formula_expected": fr[5].value == f"=F$3*(($D{index}*0.001)/$C{index})*100",
                "omni_formula_expected": fr[6].value == f"=G$3*(($D{index}*0.001)/$C{index})*100",
                "cell_evidence_json": cell_evidence(fr, cr, 9),
            }
            for label in ("fine", "omni"):
                computed, cached = row[f"{label}_ws_percent"], row[f"{label}_cached_percent"]
                row[f"{label}_cache_difference_pp"] = (
                    abs(computed - cached) if computed is not None and cached is not None else None
                )
            rows.append(row)
    return rows


def parse_summary(formula_book, cached_book):
    fs, cs = formula_book["Sheet 1"], cached_book["Sheet 1"]
    site, season, header_row = None, None, None
    rows = []
    for fr, cr in _paired_rows(fs, cs):
        first = cr[0].value
        if isinstance(first, str) and first.startswith("Fairbanks "):
            heading, season = first.removeprefix("Fairbanks ").split(", ")
            site, header_row = SITES[heading], cr[0].row
            if (cs.cell(header_row + 2, 9).value, cs.cell(header_row + 2, 10).value) != ("max", "min"):
                raise ValueError("Unexpected radiocarbon-bound order")
        day, kind = date_value(first)
        if day is None:
            continue
        if site is None:
            raise ValueError("Unassigned summary site block")
        index = cr[0].row
        high, low = number(cr[8].value), number(cr[9].value)
        if low is not None and high is not None and low > high:
            raise ValueError("Reversed source radiocarbon interval")
        row = {
            "site": site, "season": season, "date": day, "date_kind": kind,
            "workbook": "Fairbanks summary 08-11_cpp.xlsx", "sheet": "Sheet 1",
            "source_range": f"A{index}:J{index}", "site_header_cell": f"A{header_row}",
            "radiocarbon_low_percent": low, "radiocarbon_high_percent": high,
            "radiocarbon_low_flag": value_flag(cr[9].value),
            "radiocarbon_high_flag": value_flag(cr[8].value),
            "radiocarbon_interval_present": low is not None and high is not None,
            "cell_evidence_json": cell_evidence(fr, cr, 10),
        }
        for label, col in (("pm25", 1), ("pm25_uncertainty", 2), ("oc", 3),
                           ("oc_uncertainty", 4), ("ec", 5), ("ec_uncertainty", 6)):
            row[f"{label}_ug_m3"] = number(cr[col].value)
            row[f"{label}_flag"] = value_flag(cr[col].value)
        row["levoglucosan_ng_m3"] = number(cr[7].value)
        row["levoglucosan_flag"] = value_flag(cr[7].value)
        rows.append(row)
    return rows


def load_received(input_zip):
    import openpyxl

    blob = Path(input_zip).read_bytes()
    if hashlib.sha256(blob).hexdigest() != EXPECTED_ZIP_SHA256:
        raise ValueError("Input ZIP hash differs from the registered private receipt")
    with zipfile.ZipFile(BytesIO(blob)) as archive:
        if sorted(archive.namelist()) != sorted(EXPECTED_WORKBOOKS):
            raise ValueError("Unexpected or duplicate input ZIP members")
        if archive.testzip() is not None:
            raise ValueError("Input ZIP CRC failure")
        results = {}
        for name, digest in EXPECTED_WORKBOOKS.items():
            content = archive.read(name)
            if hashlib.sha256(content).hexdigest() != digest:
                raise ValueError(f"Workbook hash differs: {name}")
            # openpyxl warns about an unsupported extension being removed on save.
            # There is no save: exact originals remain untouched.
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Unknown extension is not supported.*")
                fb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
                cb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
                try:
                    results[name] = (parse_summary if name.startswith("Fairbanks") else parse_levoglucosan)(fb, cb)
                finally:
                    fb.close()
                    cb.close()
        return results["LevoglucosanResults_final1.xlsx"], results["Fairbanks summary 08-11_cpp.xlsx"]


def load_cmb(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def cmb_pairs(cmb_rows):
    keyed = {}
    for row in cmb_rows:
        if row["system"] not in ("epa", "omni"):
            continue  # Revised OMNI is not the prespecified comparison.
        key = (row["site"], row["date"])
        systems = keyed.setdefault(key, {})
        if row["system"] in systems:
            raise ValueError("Duplicate CMB site/date/system")
        systems[row["system"]] = row
    return keyed


def unique_index(rows):
    result = {}
    for row in rows:
        key = (row["site"], row["date"])
        if key in result:
            raise ValueError("Duplicate source site/date: manual identity review required")
        result[key] = row
    return result


def compare_markers(marker_rows, pairs, marker, *, recompute_missing=False, recompute_partial=False,
                    rounded_omni_factor=False, denominator="appendix_c"):
    if denominator not in ("appendix_c", "marker_workbook"):
        raise ValueError("Unknown PM denominator")
    eligible = [r for r in marker_rows if r["date_kind"] == "excel_date"]
    unique_index(eligible)
    comparisons, excluded = [], []
    for row in eligible:
        pair = pairs.get((row["site"], row["date"]))
        bounds = ((row["fine_cached_percent"], row["omni_cached_percent"])
                  if marker == "levoglucosan" else
                  (row["radiocarbon_low_percent"], row["radiocarbon_high_percent"]))
        native_present = None not in bounds
        native_partial = any(value is not None for value in bounds) and not native_present
        if marker == "levoglucosan" and (recompute_missing or (recompute_partial and native_partial)
                                          or (native_present and rounded_omni_factor)):
            bracket = marker_bracket(row["pm25_ug_m3"], row["levoglucosan_ng_m3"],
                                     FINE_CF, Decimal("13.3") if rounded_omni_factor else OMNI_CF)
            computed_bounds = bracket if bracket is not None else (None, None)
            if recompute_partial and native_partial and not recompute_missing:
                # Fill only the absent endpoint; preserve the received cached value.
                bounds = tuple(native if native is not None else computed
                               for native, computed in zip(bounds, computed_bounds))
            else:
                bounds = computed_bounds
        reason = None
        if None in bounds:
            reason = "source_marker_outputs_missing_or_flagged"
        elif marker == "levoglucosan" and marker_bracket(row["pm25_ug_m3"], row["levoglucosan_ng_m3"]) is None:
            reason = "marker_nonpositive_or_missing_input"
        elif pair is None or set(pair) != {"epa", "omni"}:
            reason = "no_exact_original_CMB_pair"
        elif any(r["cmb_row_status"] != "valid" for r in pair.values()):
            reason = "at_least_one_CMB_system_invalid"
        elif denominator == "marker_workbook" and (row["pm25_ug_m3"] is None or row["pm25_ug_m3"] <= 0):
            reason = "marker_workbook_denominator_missing_or_nonpositive"
        if reason:
            excluded.append({"site": row["site"], "date": row["date"], "reason": reason,
                             "sheet": row["sheet"], "source_range": row["source_range"]})
            continue
        result = {"site": row["site"], "date": row["date"], "marker": marker,
                  "workbook": row["workbook"], "sheet": row["sheet"],
                  "source_range": row["source_range"],
                  "marker_low_percent": bounds[0], "marker_high_percent": bounds[1],
                  "source_marker_outputs_present": native_present,
                  "CMB_denominator": denominator,
                  "reference_type": "conversion_factor_bracket_not_CI" if marker == "levoglucosan" else "reported_derived_woodsmoke_interval",
                  "marker_pm25_ug_m3": row["pm25_ug_m3"],
                  "filter_id_present": row.get("filter_id_present", False),
                  "cell_evidence_json": row["cell_evidence_json"]}
        for system, cmb in pair.items():
            mass, ws = Decimal(cmb["pm25_mass"]), Decimal(cmb["wood_smoke"])
            if mass <= 0:
                raise ValueError("Nonpositive PM in valid CMB row")
            denominator_mass = mass if denominator == "appendix_c" else row["pm25_ug_m3"]
            percent = 100 * ws / denominator_mass
            result[f"{system}_pm25_ug_m3"] = mass
            result[f"{system}_wood_smoke_ug_m3"] = ws
            result[f"{system}_wood_smoke_se_ug_m3"] = cmb["wood_smoke_se"]
            result[f"{system}_wood_smoke_percent"] = percent
            result[f"{system}_report_pages"] = cmb["report_pages"]
            result[f"{system}_mass_absolute_difference_ug_m3"] = (
                abs(mass - row["pm25_ug_m3"]) if row["pm25_ug_m3"] is not None else None
            )
            result.update({f"{system}_{k}": v for k, v in interval_metrics(percent, *bounds).items()})
        epa, omni = (result[f"{s}_interval_distance_pp"] for s in ("epa", "omni"))
        result["closer_to_interval"] = "epa" if epa < omni else "omni" if omni < epa else "tie"
        comparisons.append(result)
    return comparisons, excluded


def rounded(value, digits=6):
    return format(value, f".{digits}f") if value is not None else None


def summarize_comparisons(rows):
    result = {"n": len(rows), "by_site": dict(sorted(Counter(r["site"] for r in rows).items())),
              "closer_to_interval": {s: sum(r["closer_to_interval"] == s for r in rows) for s in ("epa", "omni", "tie")},
              "filter_ids_present": sum(r["filter_id_present"] for r in rows)}
    for system in ("epa", "omni"):
        mass_diff = [r[f"{system}_mass_absolute_difference_ug_m3"] for r in rows
                     if r[f"{system}_mass_absolute_difference_ug_m3"] is not None]
        result[system] = {
            "relation_counts": {label: sum(r[f"{system}_relation"] == label for r in rows)
                                for label in ("inside", "above", "below")},
            "mass_comparisons_n": len(mass_diff),
            "mass_absolute_difference_max_ug_m3": rounded(max(mass_diff)) if mass_diff else None,
            "mass_difference_le_0_04_n": sum(v <= Decimal("0.04") for v in mass_diff),
        }
        for metric in ("interval_distance_pp", "midpoint_signed_difference_pp", "midpoint_absolute_difference_pp"):
            avg = mean(r[f"{system}_{metric}"] for r in rows) if rows else None
            result[system][f"mean_{metric}"] = rounded(avg)
            result[system][f"mean_{metric}_2dp"] = rounded(avg, 2)
    return result


def marker_overlaps(levo, radiocarbon):
    indexed = unique_index(radiocarbon)
    results = []
    for row in levo:
        other = indexed.get((row["site"], row["date"]))
        if other is None:
            continue
        low = max(row["marker_low_percent"], other["marker_low_percent"])
        high = min(row["marker_high_percent"], other["marker_high_percent"])
        results.append({"site": row["site"], "date": row["date"],
                        "intervals_overlap": low <= high,
                        "gap_pp": max(low - high, Decimal(0)),
                        "levo_range": row["source_range"], "radiocarbon_range": other["source_range"]})
    return results


def flag_counts(rows):
    counts = Counter()
    for row in rows:
        for key, value in row.items():
            if key.endswith("_flag"):
                counts[f"{key}:{value}"] += 1
    return dict(sorted(counts.items()))


def make_summary(levo_rows, summary_rows, cmb_rows, levo, rc, overlaps, levo_excluded, rc_excluded):
    native_levo = [r for r in levo_rows if r["date_kind"] == "excel_date"]
    native_summary = [r for r in summary_rows if r["date_kind"] == "excel_date"]
    pairs = cmb_pairs(cmb_rows)
    cache_diffs = [r[f"{s}_cache_difference_pp"] for r in levo_rows for s in ("fine", "omni")
                   if r[f"{s}_cache_difference_pp"] is not None]
    summary = {
        "schema_version": 1,
        "decision": {"secondary_marker_concordance": "KEEP", "profile_choice_external_validation": "HOLD",
                     "baseline_change_authorized": False},
        "source_identity": {"private_zip_sha256": EXPECTED_ZIP_SHA256, "workbooks": EXPECTED_WORKBOOKS},
        "methods": {
            "join": "Exact normalized site and native Excel calendar date; no fuzzy date shift or imputation",
            "inclusion": "Both original EPA and OMNI rows valid and source-native cached marker bounds numeric; levoglucosan comparison additionally requires strictly positive levoglucosan and PM; radiocarbon comparison uses reported derived bounds, with strictly positive summary PM additionally required only for its alternate-denominator sensitivity",
            "cmb_percent": "100 * published wood_smoke_ug_m3 / published pm25_mass_ug_m3",
            "marker_percent": "100 * CF * 0.001 * levoglucosan_ng_m3 / measured_pm25_ug_m3",
            "conversion_factors": {"Fine": "9.01", "OMNI": "13.27", "cells": "F3:G4 in each levoglucosan site sheet"},
            "omni_13_3": "Not the numerical factor used in the primary analysis; tested separately in the explicitly labeled rounded-factor sensitivity",
            "radiocarbon_cells": "Sheet 1: I=max, J=min, % PM2.5 from woodsmoke, not raw 14C",
            "interval_distance": "max(low - estimate, estimate - high, 0)",
            "midpoint_error": "estimate - (low + high)/2; MAE averaged separately from interval distance",
            "ties": "Exact equality of unrounded interval distances",
            "uncertainty": "Published CMB standard errors retained privately, not propagated into marker brackets",
            "limits": "Fine/OMNI marker columns are conversion estimates, not CMB runs; bracket is not a confidence interval; no limits converted to zero",
            "missing_outputs": "Positive marker input does not authorize filling blank native woodsmoke-output cells; such recomputation is a separately labeled sensitivity",
            "mass_check": "Reported as diagnostic only; no mass tolerance filters are applied to achieve counts",
        },
        "inventory": {
            "levoglucosan_native_date_rows": len(native_levo),
            "levoglucosan_native_date_rows_by_site": dict(sorted(Counter(r["site"] for r in native_levo).items())),
            "levoglucosan_text_date_candidates": sum(r["date_kind"] == "text_date" for r in levo_rows),
            "summary_native_date_rows": len(native_summary),
            "summary_native_date_rows_by_site": dict(sorted(Counter(r["site"] for r in native_summary).items())),
            "radiocarbon_intervals": sum(r["radiocarbon_interval_present"] for r in native_summary),
            "radiocarbon_intervals_by_site": dict(sorted(Counter(r["site"] for r in native_summary if r["radiocarbon_interval_present"]).items())),
            "published_CMB_rows": len(cmb_rows), "original_CMB_site_date_pairs": len(pairs),
            "both_valid_original_CMB_pairs": sum(set(p) == {"epa", "omni"} and all(r["cmb_row_status"] == "valid" for r in p.values()) for p in pairs.values()),
        },
        "formula_audit": {
            "formula_expected_counts": {s: sum(r[f"{s}_formula_expected"] for r in levo_rows) for s in ("fine", "omni")},
            "formula_present_counts": {s: sum(r[f"{s}_formula_present"] for r in levo_rows) for s in ("fine", "omni")},
            "cache_comparisons": len(cache_diffs),
            "max_absolute_cache_difference_pp": rounded(max(cache_diffs), 12) if cache_diffs else None,
            "cache_difference_gt_1e_9_n": sum(v > Decimal("1e-9") for v in cache_diffs),
        },
        "source_flags": {"levoglucosan": flag_counts(levo_rows), "summary": flag_counts(summary_rows)},
        "exclusions": {"levoglucosan": dict(sorted(Counter(r["reason"] for r in levo_excluded).items())),
                       "radiocarbon": dict(sorted(Counter(r["reason"] for r in rc_excluded).items()))},
        "levoglucosan_comparison": summarize_comparisons(levo),
        "radiocarbon_comparison": summarize_comparisons(rc),
        "both_markers": {"n": len(overlaps), "brackets_overlap": sum(r["intervals_overlap"] for r in overlaps),
                         "brackets_disjoint": sum(not r["intervals_overlap"] for r in overlaps)},
        "limitations": [
            "EPA and OMNI differ in source universe/treatment; no isolated source-profile substitution",
            "Per-alternative final selectors, fit diagnostics, and source-native CMB control remain unrecovered",
            "OMNI levoglucosan factor depends partly on OMNI source experiments; no independent accuracy winner",
            "Site/date and PM checks support identity but do not replace absent historical filter IDs",
            "Sparse radiocarbon overlap and nonuniform marker preference prevent a uniform accuracy conclusion",
            "Received workbooks and detailed derived rows are private pending redistribution clarification",
        ],
    }
    return summary


def write_csv(path, rows):
    if not rows:
        return
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def private_directory(path):
    target = Path(path).resolve()
    permitted = (ROOT / "private" / "received_fairbanks").resolve()
    if target != permitted and permitted not in target.parents:
        raise ValueError("Detailed output must be below repository private/received_fairbanks")
    return target


def run_audit(input_zip, output_dir, public_summary=None, cmb_csv=None):
    target = private_directory(output_dir)
    cmb_path = Path(cmb_csv) if cmb_csv else ROOT / "outputs/cycle02/appendix_c_2008_2009_daily.csv"
    levo_rows, summary_rows = load_received(input_zip)
    cmb_rows = load_cmb(cmb_path)
    pairs = cmb_pairs(cmb_rows)
    levo, levo_excluded = compare_markers(levo_rows, pairs, "levoglucosan")
    rc, rc_excluded = compare_markers(summary_rows, pairs, "radiocarbon")
    overlaps = marker_overlaps(levo, rc)
    summary = make_summary(levo_rows, summary_rows, cmb_rows, levo, rc, overlaps, levo_excluded, rc_excluded)
    recomputed, _ = compare_markers(levo_rows, pairs, "levoglucosan", recompute_missing=True)
    partial, _ = compare_markers(levo_rows, pairs, "levoglucosan", recompute_partial=True)
    levo_native_mass, _ = compare_markers(levo_rows, pairs, "levoglucosan", denominator="marker_workbook")
    rc_native_mass, _ = compare_markers(summary_rows, pairs, "radiocarbon", denominator="marker_workbook")
    rounded_factor, _ = compare_markers(levo_rows, pairs, "levoglucosan", rounded_omni_factor=True)
    summary["sensitivity_analyses"] = {
        "levoglucosan_recompute_all_positive_inputs": {
            "meaning": "Derive missing woodsmoke bounds arithmetically; not the received native-output dataset",
            "newly_computed_native_missing_output_rows": sum(not r["source_marker_outputs_present"] for r in recomputed),
            "newly_computed_rows_lower_bound_above_100_percent": sum(not r["source_marker_outputs_present"] and r["marker_low_percent"] > 100 for r in recomputed),
            "comparison": summarize_comparisons(recomputed),
        },
        "levoglucosan_CMB_normalized_by_marker_workbook_PM": summarize_comparisons(levo_native_mass),
        "levoglucosan_recompute_partially_present_brackets_only": {
            "meaning": "Fill a missing native endpoint only where the other endpoint exists; exclude wholly absent native bounds; diagnostic reconstruction of the historical 53-row statement, not primary",
            "newly_computed_native_missing_output_rows": sum(not r["source_marker_outputs_present"] for r in partial),
            "comparison": summarize_comparisons(partial),
        },
        "radiocarbon_CMB_normalized_by_summary_workbook_PM": summarize_comparisons(rc_native_mass),
        "levoglucosan_rounded_OMNI_CF_13_3_native_complete_rows": summarize_comparisons(rounded_factor),
    }
    summary["source_identity"]["published_CMB_csv"] = {
        "filename": cmb_path.name, "sha256": hashlib.sha256(cmb_path.read_bytes()).hexdigest()
    }
    target.mkdir(parents=True, exist_ok=True)
    for name, rows in (("levoglucosan_rows", levo_rows), ("summary_rows", summary_rows),
                       ("levoglucosan_comparisons", levo), ("radiocarbon_comparisons", rc),
                       ("both_marker_overlap", overlaps), ("levoglucosan_exclusions", levo_excluded),
                       ("radiocarbon_exclusions", rc_excluded),
                       ("sensitivity_levo_recomputed", recomputed),
                       ("sensitivity_levo_partial_brackets", partial),
                       ("sensitivity_levo_native_PM", levo_native_mass),
                       ("sensitivity_radiocarbon_native_PM", rc_native_mass),
                       ("sensitivity_levo_CF_13_3", rounded_factor)):
        write_csv(target / f"{name}.csv", rows)
    text = json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    (target / "aggregate_summary.json").write_text(text, encoding="utf-8")
    if public_summary is not None:
        destination = Path(public_summary).resolve()
        if destination == Path(input_zip).resolve() or destination == cmb_path.resolve():
            raise ValueError("Refusing to overwrite source with public summary")
        if destination.suffix.lower() != ".json":
            raise ValueError("Public summary must have a .json extension")
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(text, encoding="utf-8")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-zip", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--public-summary", type=Path)
    parser.add_argument("--cmb-csv", type=Path, help="Defaults to frozen Cycle 02 Appendix C CSV")
    args = parser.parse_args()
    result = run_audit(args.input_zip, args.output_dir, args.public_summary, args.cmb_csv)
    print(json.dumps({"inventory": result["inventory"],
                      "levoglucosan_comparison": result["levoglucosan_comparison"],
                      "radiocarbon_comparison": result["radiocarbon_comparison"],
                      "both_markers": result["both_markers"]}, indent=2))


if __name__ == "__main__":
    main()
