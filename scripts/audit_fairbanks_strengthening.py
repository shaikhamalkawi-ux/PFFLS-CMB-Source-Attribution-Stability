#!/usr/bin/env python3
"""Read-only audit of privately supplied Fairbanks workbooks.

Public outputs are aggregate only. Per-row data and cell diagnostics are written
only beneath the gitignored private directory. No workbook is saved or edited.
The original spreadsheets are required for the private integration tests.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
from statistics import mean
import warnings

import openpyxl

REPO = Path(__file__).resolve().parents[1]
INPUT_DIR = REPO.parent / "_inputs/strengthening_20260926/fairbanks"
PRIVATE_DIR = REPO / "private/strengthening_20260926/fairbanks"
HASHES = {
    "Fairbanks summary 08-11_cpp.xlsx": "973917d6c504f25980938a8a45da62ce766099734e053a65882c2c05391d8ee7",
    "LevoglucosanResults_final1.xlsx": "1541b1102f49fec9d461fbfb9c3095c13d7b8a23cbce54c81f1744fe7572691a",
}
CMB_LF_SHA256 = "9c7d225b394b19ba9354077a50bf24c3f77a3972064c42e8e090168da749b172"
SITE_NAMES = {"State Building": "state_building", "Peger Rd": "peger_road", "North Pole": "north_pole", "RAMS": "rams"}


def numeric(value):
    return isinstance(value, (float, int)) and not isinstance(value, bool) and math.isfinite(value)


def value_flag(value):
    if value is None:
        return "missing"
    if numeric(value):
        return "negative" if value < 0 else "zero" if value == 0 else "positive"
    return "source_text_flag"


def parse_date(value):
    """No inferred serial dates or mixed ID/date strings. Explicit US text dates."""
    if isinstance(value, datetime):
        return value.date().isoformat(), "excel_date"
    if isinstance(value, date):
        return value.isoformat(), "excel_date"
    if isinstance(value, str) and re.fullmatch(r"\s*\d{1,2}/\d{1,2}/\d{4}\s*", value):
        return datetime.strptime(value.strip(), "%m/%d/%Y").date().isoformat(), "us_text_date"
    return None, None


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def load_pair(path):
    # openpyxl warns about unsupported extensions on read. No save is performed.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Unknown extension is not supported")
        formula = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return ({s.title: list(s.iter_rows(values_only=True)) for s in formula},
                    {s.title: list(s.iter_rows(values_only=True)) for s in cached})
        finally:
            formula.close()
            cached.close()


def interval_distance(value, lower, upper):
    if not all(numeric(v) for v in (value, lower, upper)) or lower > upper:
        raise ValueError("invalid interval")
    return max(lower - value, 0, value - upper)


def inventory_duplicates(rows):
    counts = Counter((r["site"], r["date"]) for r in rows)
    return [{"site": s, "date": d, "count": n} for (s, d), n in counts.items() if n > 1]


def extract_levo(formula, cached):
    rows, diagnostics, exclusions, cache_checks = [], [], [], []
    formula_errors = []
    for sheet, data in cached.items():
        site = SITE_NAMES[sheet]
        cf_low, cf_high = data[2][5:7]
        assert numeric(cf_low) and numeric(cf_high) and cf_low < cf_high
        for index, values in enumerate(data, 1):
            for col, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NAME?", "#NULL!")):
                    formula_errors.append({"sheet": sheet, "row": index, "column": col, "error": value})
            dt, dt_type = parse_date(values[0])
            if dt is None:
                continue
            pm, concentration, share, low, high = values[2:7]
            row = {"site": site, "date": dt, "date_type": dt_type, "sheet": sheet,
                   "row": index, "filter_id": values[1], "pm25_ug_m3": pm,
                   "levo_ng_m3": concentration, "share_cached_percent": share,
                   "low_cached_percent": low, "high_cached_percent": high,
                   "cf_low": cf_low, "cf_high": cf_high}
            if numeric(pm) and pm > 0 and numeric(concentration):
                recomputed_share = 100 * concentration * 0.001 / pm
                row.update(share_recomputed_percent=recomputed_share,
                           low_recomputed_percent=cf_low * recomputed_share,
                           high_recomputed_percent=cf_high * recomputed_share)
                for col_name, col_index in [("share", 4), ("low", 5), ("high", 6)]:
                    actual = values[col_index]
                    recomputed = row[f"{col_name}_recomputed_percent"]
                    source_formula = formula[sheet][index - 1][col_index]
                    if isinstance(source_formula, str) and source_formula.startswith("="):
                        cache_checks.append({"site": site, "date": dt, "sheet": sheet,
                                             "row": index, "column": col_index + 1,
                                             "column_name": col_name,
                                             "agrees": numeric(actual) and abs(actual - recomputed) <= 1e-8,
                                             "absolute_difference": abs(actual - recomputed) if numeric(actual) else None})
                    if not numeric(actual) or abs(actual - recomputed) > 1e-8:
                        diagnostics.append({"site": site, "date": dt, "sheet": sheet,
                                            "row": index, "column": col_index + 1,
                                            "formula": source_formula,
                                            "issue": "missing_output_not_formula_error" if source_formula is None else "cached_value_mismatch",
                                            "cached": actual, "recomputed": recomputed,
                                            "difference": actual - recomputed if numeric(actual) else None})
            else:
                exclusions.append({"site": site, "date": dt, "reason": "nonpositive_or_missing_PM_or_missing_levo"})
            rows.append(row)
    return rows, diagnostics, exclusions, formula_errors, cache_checks


def extract_summary(formula, cached):
    rows, annotations, errors = [], [], []
    site, season = None, None
    for index, values in enumerate(cached["Sheet 1"], 1):
        if isinstance(values[0], str) and values[0].startswith("Fairbanks "):
            label, season = values[0].removeprefix("Fairbanks ").split(", ")
            site = SITE_NAMES[label]
        dt, dt_type = parse_date(values[0])
        if dt:
            pm, pm_uncertainty, levo, upper, lower = values[1], values[2], values[7], values[8], values[9]
            rows.append({"site": site, "season": season, "date": dt,
                         "date_type": dt_type, "row": index, "pm25_ug_m3": pm,
                         "pm25_uncertainty_ug_m3": pm_uncertainty,
                         "levo_ng_m3": levo, "radio_low_percent": lower,
                         "radio_high_percent": upper})
        elif isinstance(values[0], str) and re.search(r"\d/\d+", values[0]):
            annotations.append({"row": index, "label": values[0],
                                "radio_low_percent": values[9], "radio_high_percent": values[8]})
    for sheet, data in cached.items():
        for index, values in enumerate(data, 1):
            for col, value in enumerate(values, 1):
                if isinstance(value, str) and value.startswith(("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NAME?", "#NULL!")):
                    errors.append({"sheet": sheet, "row": index, "column": col, "error": value,
                                   "formula": formula[sheet][index-1][col-1]})
    return rows, annotations, errors


def load_cmb(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    lookup = {}
    for row in rows:
        key = row["site"], row["date"], row["system"]
        if key in lookup:
            raise ValueError("duplicate public Appendix C key")
        lookup[key] = row
    return lookup


def pair_values(lookup, site, dt):
    values = [lookup.get((site, dt, system)) for system in ("epa", "omni")]
    if not all(v and v["cmb_row_status"] == "valid" for v in values):
        return None
    epa, omni = values
    masses = [float(v["pm25_mass"]) for v in values]
    if min(masses) <= 0 or masses[0] != masses[1]:
        raise ValueError("CMB PM mass inconsistency")
    return {"cmb_pm25_ug_m3": masses[0],
            "epa_percent": 100 * float(epa["wood_smoke"]) / masses[0],
            "omni_percent": 100 * float(omni["wood_smoke"]) / masses[1],
            "epa_wood_smoke_se_ug_m3": float(epa["wood_smoke_se"]),
            "omni_wood_smoke_se_ug_m3": float(omni["wood_smoke_se"])}


def annotate_distances(row, low_key, high_key):
    low, high = row[low_key], row[high_key]
    for system in ("epa", "omni"):
        value = row[f"{system}_percent"]
        row[f"{system}_distance_pp"] = interval_distance(value, low, high)
        row[f"{system}_position"] = "below" if value < low else "above" if value > high else "inside"
        row[f"{system}_midpoint_abs_pp"] = abs(value - (low + high) / 2)
    difference = row["epa_distance_pp"] - row["omni_distance_pp"]
    row["closer_system"] = "tie" if abs(difference) < 1e-10 else "omni" if difference > 0 else "epa"
    return row


def aggregate_pairs(rows):
    if not rows:
        return {"n": 0}
    result = {"n": len(rows), "sites": dict(sorted(Counter(r["site"] for r in rows).items())),
              "closer_system": dict(sorted(Counter(r["closer_system"] for r in rows).items()))}
    for system in ("epa", "omni"):
        result[system] = {"positions": dict(sorted(Counter(r[f"{system}_position"] for r in rows).items())),
                          "mean_interval_distance_pp": mean(r[f"{system}_distance_pp"] for r in rows),
                          "midpoint_mae_pp": mean(r[f"{system}_midpoint_abs_pp"] for r in rows)}
    return result


def write_csv(path, rows):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields = list(dict.fromkeys(k for row in rows for k in row))
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def run_audit(input_dir=INPUT_DIR, private_dir=PRIVATE_DIR):
    input_dir, private_dir = Path(input_dir), Path(private_dir)
    allowed = REPO / "private"
    if not private_dir.resolve().is_relative_to(allowed.resolve()):
        raise ValueError("Per-row evidence may be written only inside repository/private")
    for name, expected in HASHES.items():
        if sha256(input_dir / name) != expected:
            raise ValueError(f"Unexpected source hash: {name}")
    lf, lc = load_pair(input_dir / "LevoglucosanResults_final1.xlsx")
    sf, sc = load_pair(input_dir / "Fairbanks summary 08-11_cpp.xlsx")
    levo, formula_diagnostics, levo_invalid, levo_errors, cache_checks = extract_levo(lf, lc)
    summary, annotations, summary_errors = extract_summary(sf, sc)
    for data in (levo, summary):
        if inventory_duplicates(data):
            raise ValueError("duplicate site/date keys require explicit resolution")
    cmb_path = REPO / "outputs/cycle02/appendix_c_2008_2009_daily.csv"
    cmb_lf_sha = hashlib.sha256(cmb_path.read_bytes().replace(b"\r\n", b"\n")).hexdigest()
    if cmb_lf_sha != CMB_LF_SHA256:
        raise ValueError("Unexpected public Appendix C snapshot hash")
    lookup = load_cmb(cmb_path)
    levo_pairs, recomputed_pairs, partial_pairs, levo_excluded, radio_pairs = [], [], [], [], []
    for row in levo:
        paired = pair_values(lookup, row["site"], row["date"])
        if paired is None:
            continue
        base = dict(row, **paired)
        base["pm_mass_abs_difference_ug_m3"] = abs(row["pm25_ug_m3"] - paired["cmb_pm25_ug_m3"]) if numeric(row["pm25_ug_m3"]) else None
        if all(numeric(row[k]) and row[k] > 0 for k in ("low_cached_percent", "high_cached_percent")):
            levo_pairs.append(annotate_distances(base, "low_cached_percent", "high_cached_percent"))
        else:
            levo_excluded.append(dict(base, reason="not_two_positive_cached_marker_endpoints"))
        if all(numeric(row.get(k)) and row[k] > 0 for k in ("low_recomputed_percent", "high_recomputed_percent")):
            computed = annotate_distances(dict(row, **paired), "low_recomputed_percent", "high_recomputed_percent")
            recomputed_pairs.append(computed)
            if any(numeric(row[k]) for k in ("low_cached_percent", "high_cached_percent")):
                partial_pairs.append(computed)
    for row in summary:
        paired = pair_values(lookup, row["site"], row["date"])
        if paired and all(numeric(row[k]) and row[k] > 0 for k in ("radio_low_percent", "radio_high_percent")):
            radio_pairs.append(annotate_distances(dict(row, **paired), "radio_low_percent", "radio_high_percent"))
    levo_lookup = {(r["site"], r["date"]): r for r in levo_pairs}
    overlap = []
    for row in radio_pairs:
        other = levo_lookup.get((row["site"], row["date"]))
        if other:
            overlap.append({"site": row["site"], "date": row["date"],
                            "intervals_overlap": max(row["radio_low_percent"], other["low_cached_percent"]) <= min(row["radio_high_percent"], other["high_cached_percent"])})
    matched_keys = {(r["site"], r["date"]) for r in levo_pairs}
    radio_native_pm = []
    partial_native_pm = []
    for pairs, target, low_key, high_key in [
        (radio_pairs, radio_native_pm, "radio_low_percent", "radio_high_percent"),
        (partial_pairs, partial_native_pm, "low_recomputed_percent", "high_recomputed_percent"),
    ]:
        for row in pairs:
            if numeric(row["pm25_ug_m3"]) and row["pm25_ug_m3"] > 0:
                adjusted = dict(row)
                for system in ("epa", "omni"):
                    adjusted[f"{system}_percent"] *= row["cmb_pm25_ug_m3"] / row["pm25_ug_m3"]
                target.append(annotate_distances(adjusted, low_key, high_key))
    result = {
        "source_sha256": HASHES,
        "public_cmb_csv_sha256": sha256(cmb_path),
        "public_cmb_csv_lf_normalized_sha256": cmb_lf_sha,
        "levo_inventory": {"dated_rows": len(levo), "by_site": dict(Counter(r["site"] for r in levo)),
                           "native_excel_date_rows": sum(r["date_type"] == "excel_date" for r in levo),
                           "native_excel_rows_by_site": dict(Counter(r["site"] for r in levo if r["date_type"] == "excel_date")),
                           "date_types": dict(Counter(r["date_type"] for r in levo)),
                           "duplicate_site_date_keys": len(inventory_duplicates(levo)),
                           "rows_with_filter_id": sum(r["filter_id"] is not None for r in levo)},
        "summary_inventory": {"dated_rows": len(summary), "by_site": dict(Counter(r["site"] for r in summary)),
                              "duplicate_site_date_keys": len(inventory_duplicates(summary)),
                              "dated_rows_with_two_radio_bounds": sum(all(numeric(r[k]) for k in ("radio_low_percent", "radio_high_percent")) for r in summary),
                              "mixed_id_site_date_annotations_excluded": len(annotations)},
        "cached_formula_audit": {"existing_formula_checks": len(cache_checks),
                                 "existing_endpoint_formula_checks": sum(r["column_name"] in ("low", "high") for r in cache_checks),
                                 "existing_formula_bad_caches": sum(not r["agrees"] for r in cache_checks),
                                 "missing_source_output_cells": sum(r["issue"] == "missing_output_not_formula_error" for r in formula_diagnostics),
                                 "rows_with_missing_source_outputs": len({(r["site"], r["date"]) for r in formula_diagnostics}),
                                 "primary_rows_with_missing_source_outputs": len({(r["site"], r["date"]) for r in formula_diagnostics} & matched_keys),
                                 "levo_excel_errors": len(levo_errors), "summary_excel_errors": len(summary_errors)},
        "source_flags": {
            "levo_pm25": dict(Counter(value_flag(r["pm25_ug_m3"]) for r in levo)),
            "levo_concentration": dict(Counter(value_flag(r["levo_ng_m3"]) for r in levo)),
            "summary_pm25": dict(Counter(value_flag(r["pm25_ug_m3"]) for r in summary)),
            "radio_upper": dict(Counter(value_flag(r["radio_high_percent"]) for r in summary)),
        },
        "levo_cached": aggregate_pairs(levo_pairs),
        "levo_recomputed": aggregate_pairs(recomputed_pairs),
        "levo_partial_recomputed_sensitivity": aggregate_pairs(partial_pairs),
        "levo_partial_native_pm_sensitivity": aggregate_pairs(partial_native_pm),
        "radiocarbon": aggregate_pairs(radio_pairs),
        "radiocarbon_native_pm_sensitivity": aggregate_pairs(radio_native_pm),
        "paired_levo_rows_without_filter_id": sum(r["filter_id"] is None for r in levo_pairs),
        "maximum_paired_levo_pm_difference_ug_m3": max(r["pm_mass_abs_difference_ug_m3"] for r in levo_pairs),
        "both_marker_dates": len(overlap), "marker_overlap_dates": sum(r["intervals_overlap"] for r in overlap),
        "admission": {"secondary_historical_system_concordance": "KEEP", "external_profile_choice_accuracy": "HOLD"},
    }
    private_dir.mkdir(parents=True, exist_ok=True)
    for name, rows in {"levoglucosan_rows": levo, "summary_rows": summary,
                       "formula_diagnostics": formula_diagnostics, "levo_invalid_inputs": levo_invalid,
                       "existing_formula_cache_checks": cache_checks,
                       "levo_excel_errors": levo_errors, "summary_excel_errors": summary_errors,
                       "summary_mixed_annotations": annotations, "levoglucosan_comparisons": levo_pairs,
                       "levoglucosan_recomputed_comparisons": recomputed_pairs,
                       "levoglucosan_partial_recomputed_comparisons": partial_pairs,
                       "levoglucosan_partial_native_pm_comparisons": partial_native_pm,
                       "levoglucosan_exclusions": levo_excluded, "radiocarbon_comparisons": radio_pairs,
                       "radiocarbon_native_pm_comparisons": radio_native_pm,
                       "marker_overlap": overlap}.items():
        write_csv(private_dir / f"{name}.csv", rows)
    (private_dir / "aggregate_summary.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    for name, expected in HASHES.items():
        assert sha256(input_dir / name) == expected, "source changed during read-only audit"
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=INPUT_DIR)
    parser.add_argument("--private-output-dir", type=Path, default=PRIVATE_DIR)
    args = parser.parse_args()
    print(json.dumps(run_audit(args.input_dir, args.private_output_dir), indent=2))


if __name__ == "__main__":
    main()
